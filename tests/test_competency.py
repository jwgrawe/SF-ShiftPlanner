"""Competency matrix, Excel in/out with review and three-way merge, name
linking, the temps' standard set (D96–D99) and manual shift times (D95)."""
import datetime as dt
import io
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import hr_fixtures as fx  # noqa: E402
from app import competency, db, hr_import, importer  # noqa: E402
from app.hr_export import Rejected  # noqa: E402

D = dt.date
START = D(2026, 12, 28)
EXPORTED = dt.datetime(2026, 12, 10, 12, 40)
TODAY = D(2027, 1, 4)
HEADER = ["Navn på ansatt", "Produksjon, uren sone", "Kontrollsone", "Arbeidsbord/brikkelegging",
          "Sterrad + poliklinikker/løspakk"]


def workbook(rows, header=HEADER) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append([None, "Uren sone"])
    ws.append(header)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class MarksAndNamesTest(unittest.TestCase):
    def test_marks(self):
        for value, status in (("Ja", "qualified"), ("x", "qualified"), ("Ukjent", "uncertain"),
                              ("?", "uncertain"), ("", ""), (None, ""), ("Nei", "")):
            self.assertEqual(competency.parse_mark(value), (status, True), value)
        self.assertEqual(competency.parse_mark("kanskje"), (None, False))

    def test_abbreviated_names_find_their_person(self):
        people = ["Hansen Kari Marie", "Hansen Kjell", "Gråwe-Kjer Jo Willy", "Olsen Tor",
                  "Berg Anne", "Bergersen Anne-Lise"]
        aliases = ["Kari H.", "Hansen, Kj.", "Jo Gråwe-Kjer", "Tor O", "Anne Berg", "A-L Bergersen"]
        result = competency.link_suggestions(aliases, people)
        self.assertEqual(result["Kari H."]["choice"], "Hansen Kari Marie")
        self.assertEqual(result["Hansen, Kj."]["choice"], "Hansen Kjell")
        self.assertEqual(result["Jo Gråwe-Kjer"]["choice"], "Gråwe-Kjer Jo Willy")
        self.assertEqual(result["Tor O"]["choice"], "Olsen Tor")
        self.assertEqual(result["Anne Berg"]["choice"], "Berg Anne")
        self.assertEqual(result["Jo Gråwe-Kjer"]["confidence"], "sikker")
        chosen = [r["choice"] for r in result.values() if r["choice"]]
        self.assertEqual(len(chosen), len(set(chosen)))          # one-to-one

    def test_no_suggestion_for_a_stranger(self):
        result = competency.link_suggestions(["Zeta Ukjent"], ["Hansen Kari", "Olsen Tor"])
        self.assertEqual(result["Zeta Ukjent"]["confidence"], "ingen")


class RealModeBase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.inbox = self.tmp / "import"
        self.inbox.mkdir()
        db.set_mode("real")
        self.conn = db.get_conn(self.tmp / "t.db")
        db.init_schema(self.conn)
        self.conn.execute("PRAGMA foreign_keys = OFF")
        with self.conn:
            importer.import_seed(self.conn, people=False)
        self.conn.execute("PRAGMA foreign_keys = ON")
        people = [("Hansen Kari", "Sterilforsyningstekniker", {START + dt.timedelta(d): "D" for d in range(28)}),
                  ("Olsen Tor", "Sterilforsyningstekniker", {START + dt.timedelta(d): "A" for d in range(28)})]
        self.drop(fx.plan_name(EXPORTED), fx.plan_csv(people, START, 4, EXPORTED))
        week = fx.ukeplan_csv([("1001", "Hansen Kari", {TODAY: "D"}),
                               ("2001", "Vikar Vera", {TODAY: "D", TODAY + dt.timedelta(1): "DX"})],
                              TODAY, dt.datetime(2027, 1, 4, 6))
        self.drop(fx.ukeplan_name(dt.datetime(2027, 1, 4, 6)), week)
        hr_import.scan_inbox(self.conn, self.inbox)
        hr_import.rebuild(self.conn, TODAY)

    def tearDown(self):
        self.conn.close()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def drop(self, name, raw):
        path = self.inbox / name
        path.write_bytes(raw)
        os.utime(path, (0, 0))

    def comps(self, full_name):
        return {r["competency_id"]: (r["status"], r["source"]) for r in self.conn.execute(
            """SELECT ec.* FROM employee_competencies ec JOIN employees e USING (employee_id)
               WHERE e.full_name = ?""", (full_name,))}


class MatrixTest(RealModeBase):
    def test_matrix_lists_people_with_temps_flagged(self):
        m = competency.matrix(self.conn)
        rows = {r["full_name"]: r for r in m["rows"]}
        self.assertEqual(set(rows), {"Hansen Kari", "Olsen Tor", "Vikar Vera"})
        self.assertTrue(rows["Vikar Vera"]["temp"])
        self.assertFalse(rows["Hansen Kari"]["temp"])
        self.assertEqual(rows["Hansen Kari"]["title"], "Sterilforsyningstekniker")
        combined = [c for c in m["columns"] if c["key"].startswith("Sterrad")][0]
        self.assertEqual(combined["ids"], ["sterrad", "poliklinikker_lospakk"])

    def test_edits_save_and_stale_edits_are_refused(self):
        applied, conflicts = competency.apply_changes(self.conn, [
            {"key": "Hansen Kari", "col": "Kontrollsone", "old": "", "new": "qualified"}])
        self.assertEqual((applied, conflicts), (1, []))
        # Someone working from an older page still thinks the cell is empty.
        applied, conflicts = competency.apply_changes(self.conn, [
            {"key": "Hansen Kari", "col": "Kontrollsone", "old": "", "new": "uncertain"}])
        self.assertEqual(applied, 0)
        self.assertEqual(len(conflicts), 1)
        hr_import.rebuild(self.conn, TODAY)
        self.assertEqual(self.comps("Hansen Kari"), {"kontrollsone": ("qualified", None)})
        log = [r[0] for r in self.conn.execute("SELECT text FROM import_log WHERE area = 'kompetanse'")]
        self.assertTrue(any("Kontrollsone Nei → Ja" in line for line in log))

    def test_combined_column_sets_both_competencies(self):
        competency.apply_changes(self.conn, [{"key": "Hansen Kari", "col": "Sterrad + poliklinikker/løspakk",
                                              "old": "", "new": "qualified"}])
        hr_import.rebuild(self.conn, TODAY)
        self.assertEqual(set(self.comps("Hansen Kari")), {"sterrad", "poliklinikker_lospakk"})

    def test_standard_set_applies_to_temps_only(self):
        hr_import.rebuild(self.conn, TODAY)
        self.assertTrue(any("Standardkompetansen for vikarer er ikke satt" in r[0] for r in
                            self.conn.execute("SELECT text FROM import_findings")))
        competency.apply_changes(self.conn, [
            {"key": competency.STANDARD_KEY, "col": "Produksjon, uren sone", "old": "", "new": "qualified"},
            {"key": competency.STANDARD_KEY, "col": "Arbeidsbord/brikkelegging", "old": "", "new": "qualified"}])
        hr_import.rebuild(self.conn, TODAY)
        self.assertEqual(self.comps("Vikar Vera"), {"produksjon_uren": ("qualified", "standard"),
                                                    "arbeidsbord": ("qualified", "standard")})
        self.assertEqual(self.comps("Olsen Tor"), {})               # regular staff: assessed only
        # A temp with own rows uses those instead.
        competency.apply_changes(self.conn, [{"key": "Vikar Vera", "col": "Kontrollsone",
                                              "old": "", "new": "qualified"}])
        hr_import.rebuild(self.conn, TODAY)
        self.assertEqual(self.comps("Vikar Vera"), {"kontrollsone": ("qualified", None)})


class WorkbookTest(RealModeBase):
    def test_unknown_columns_and_duplicate_names_are_refused(self):
        with self.assertRaises(Rejected):
            competency.read_workbook(self.conn, workbook([["Hansen Kari", "Ja"]],
                                                         header=HEADER + ["Sjokolade"]))
        with self.assertRaises(Rejected):
            competency.read_workbook(self.conn, workbook([["Hansen Kari", "Ja"], ["Hansen Kari", ""]]))

    def test_unclear_marks_leave_the_cell_alone(self):
        data = competency.read_workbook(self.conn, workbook([["Hansen Kari", "kanskje", "Ja"]]))
        self.assertNotIn("Produksjon, uren sone", data["rows"]["Hansen Kari"])
        self.assertTrue(any("kanskje" in n for n in data["notes"]))

    def test_review_links_names_then_shows_overwrites(self):
        competency.apply_changes(self.conn, [{"key": "Olsen Tor", "col": "Kontrollsone",
                                              "old": "", "new": "qualified"}])
        with self.conn:
            import_id = competency.stage(self.conn, "liste.xlsx", workbook([
                ["Kari H.", "Ja", "Ukjent", "Ja"],
                ["Tor Olsen", "Ja", "", "Ja"],          # Kontrollsone Ja -> blank: overwrite
            ]))
        model = competency.review(self.conn, import_id)
        self.assertEqual(model["step"], "link")
        self.assertEqual(model["suggestions"]["Kari H."]["choice"], "Hansen Kari")
        self.assertEqual(model["suggestions"]["Tor Olsen"]["choice"], "Olsen Tor")
        self.assertEqual(competency.save_links(self.conn, {"Kari H.": "Hansen Kari",
                                                           "Tor Olsen": "Hansen Kari"})[0][:1], "«")
        self.assertEqual(competency.save_links(self.conn, {"Kari H.": "Hansen Kari",
                                                           "Tor Olsen": "Olsen Tor"}), [])
        model = competency.review(self.conn, import_id)
        self.assertEqual(model["step"], "diff")
        self.assertEqual(model["counts"], {"ny": 5, "overskriver": 1})
        competency.apply_import(self.conn, import_id, overwrite_conflicts=False)
        values = competency.current_values(self.conn)
        self.assertEqual(values["Olsen Tor"], {"Produksjon, uren sone": "qualified",
                                               "Arbeidsbord/brikkelegging": "qualified"})
        # Linked names are remembered: the same file again needs no linking.
        with self.conn:
            again = competency.stage(self.conn, "liste.xlsx", workbook([["Kari H.", "Ja"]]))
        self.assertEqual(competency.review(self.conn, again)["step"], "diff")

    def test_generated_form_round_trip_applies_only_what_was_changed(self):
        from openpyxl import load_workbook

        competency.apply_changes(self.conn, [
            {"key": "Hansen Kari", "col": "Kontrollsone", "old": "", "new": "qualified"},
            {"key": "Olsen Tor", "col": "Kontrollsone", "old": "", "new": "qualified"}])
        form = competency.build_form(self.conn)
        wb = load_workbook(io.BytesIO(form))
        self.assertEqual(wb[competency.BASE_SHEET].sheet_state, "veryHidden")
        ws = wb[competency.MAIN_SHEET]
        names = [ws.cell(row=r, column=1).value for r in range(3, ws.max_row + 1)]
        self.assertEqual(names, ["Hansen Kari", "Olsen Tor"])     # no temps in the form
        heads = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
        kontroll = heads.index("Kontrollsone") + 1
        arbeid = heads.index("Arbeidsbord/brikkelegging") + 1
        ws.cell(row=3, column=arbeid, value="Ja")                 # Kari: filled in the form
        ws.cell(row=4, column=kontroll, value="Ukjent")           # Tor: changed in the form …
        buffer = io.BytesIO()
        wb.save(buffer)
        # … while in the app someone removed Tor's Kontrollsone, and also
        # changed Kari's, which the form leaves untouched.
        competency.apply_changes(self.conn, [
            {"key": "Olsen Tor", "col": "Kontrollsone", "old": "qualified", "new": ""},
            {"key": "Hansen Kari", "col": "Kontrollsone", "old": "qualified", "new": "uncertain"}])
        with self.conn:
            import_id = competency.stage(self.conn, "skjema.xlsx", buffer.getvalue())
        model = competency.review(self.conn, import_id)
        self.assertTrue(model["has_base"])
        self.assertEqual(model["counts"], {"ny": 1, "konflikt": 1})
        competency.apply_import(self.conn, import_id, overwrite_conflicts=False)
        values = competency.current_values(self.conn)
        self.assertEqual(values["Hansen Kari"], {"Kontrollsone": "uncertain",          # app edit kept
                                                 "Arbeidsbord/brikkelegging": "qualified"})
        self.assertEqual(values.get("Olsen Tor", {}), {})                                # conflict: app kept


class ManualTimesTest(RealModeBase):
    def pending(self):
        return {(r["date"], r["raw"]): r["fixable"] for r in self.conn.execute(
            "SELECT * FROM roster_pending")}

    def test_entered_times_plan_the_shift_until_the_cell_changes(self):
        self.assertEqual(self.pending(), {("2027-01-05", "DX"): 1})
        vera = self.conn.execute("SELECT employee_id FROM employees WHERE full_name = 'Vikar Vera'").fetchone()[0]
        with self.conn:
            self.conn.execute("INSERT INTO manual_shift_times VALUES (?, '2027-01-05', 'DX', '07:30', '15:30', 'nå')",
                              (vera,))
        hr_import.rebuild(self.conn, TODAY)
        row = self.conn.execute("SELECT * FROM roster WHERE employee_id = ? AND date = '2027-01-05'",
                                (vera,)).fetchone()
        self.assertEqual((row["shift_code"], row["start_override"], row["end_override"],
                          row["category_override"]), ("MAN", "07:30", "15:30", "tidligvakt"))
        self.assertEqual(self.pending(), {})
        from app import planner
        self.assertEqual(planner.day_shifts(self.conn, D(2027, 1, 5))[vera]["start"], "07:30")
        # A newer ukeplan writes something else in that cell: asked again.
        week = fx.ukeplan_csv([("1001", "Hansen Kari", {TODAY: "D"}),
                               ("2001", "Vikar Vera", {TODAY + dt.timedelta(1): "DY"})],
                              TODAY, dt.datetime(2027, 1, 5, 9))
        self.drop(fx.ukeplan_name(dt.datetime(2027, 1, 5, 9)), week)
        hr_import.scan_inbox(self.conn, self.inbox)
        hr_import.rebuild(self.conn, TODAY)
        self.assertEqual(self.pending(), {("2027-01-05", "DY"): 1})

    def test_a_night_entered_by_hand_ends_on_the_shown_date(self):
        vera = self.conn.execute("SELECT employee_id FROM employees WHERE full_name = 'Vikar Vera'").fetchone()[0]
        with self.conn:
            self.conn.execute("INSERT INTO manual_shift_times VALUES (?, '2027-01-05', 'DX', '22:00', '07:00', 'nå')",
                              (vera,))
        hr_import.rebuild(self.conn, TODAY)
        # Moved back to Monday evening — where Vera already has a D shift, so
        # the night is held rather than one of the two silently dropped.
        reason = self.conn.execute(
            "SELECT reason FROM roster_pending WHERE employee_id = ? AND date = '2027-01-04'",
            (vera,)).fetchone()[0]
        self.assertEqual(reason, "to vakter samme dag")


if __name__ == "__main__":
    unittest.main()
